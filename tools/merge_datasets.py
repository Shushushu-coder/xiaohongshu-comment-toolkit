#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
营销对话 + 真实评论 数据合并脚本

将两种数据源合并成统一格式，用于大模型训练
"""

import json
from pathlib import Path
from docx import Document
from typing import List, Dict

PROJECT_ROOT = Path(__file__).resolve().parents[1]


def extract_marketing_dialogues(docx_path: str) -> List[str]:
    """从DOCX提取营销对话"""
    doc = Document(docx_path)
    
    dialogues = []
    current_dialogue = []
    
    for para in doc.paragraphs:
        text = para.text.strip()
        if not text:
            if current_dialogue:
                dialogues.append('\n'.join(current_dialogue))
                current_dialogue = []
        else:
            # 跳过标题和说明
            if text in ['营销对话数据', '对话式'] or text.startswith('大写A指代'):
                continue
            current_dialogue.append(text)
    
    if current_dialogue:
        dialogues.append('\n'.join(current_dialogue))
    
    return dialogues


def load_real_dialogues(txt_path: str) -> List[str]:
    """加载转换后的真实评论对话"""
    with open(txt_path, 'r', encoding='utf-8') as f:
        content = f.read()
    
    # 分割对话（跳过标题块：第一段是标题+分隔线，其后每段一组对话）
    parts = content.split('\n\n')
    dialogues = [part.strip() for part in parts[1:] if part.strip()]
    
    return dialogues


def create_jsonl_dataset(marketing_dialogues: List[str], 
                         real_dialogues: List[str],
                         output_path: str):
    """
    创建JSONL格式数据集
    每条数据包含：dialogue, type, source
    """
    with open(output_path, 'w', encoding='utf-8') as f:
        # 写入营销对话
        for i, dialogue in enumerate(marketing_dialogues, 1):
            data = {
                'id': f'marketing_{i}',
                'dialogue': dialogue,
                'type': 'marketing',
                'source': 'synthetic',
                'label': 0  # 0 表示营销内容
            }
            f.write(json.dumps(data, ensure_ascii=False) + '\n')
        
        # 写入真实评论
        for i, dialogue in enumerate(real_dialogues, 1):
            data = {
                'id': f'real_{i}',
                'dialogue': dialogue,
                'type': 'user_comment',
                'source': 'real',
                'label': 1  # 1 表示真实内容
            }
            f.write(json.dumps(data, ensure_ascii=False) + '\n')
    
    print(f"✅ JSONL数据集已保存到: {output_path}")
    print(f"   营销对话: {len(marketing_dialogues)} 条")
    print(f"   真实评论: {len(real_dialogues)} 条")
    print(f"   总计: {len(marketing_dialogues) + len(real_dialogues)} 条")


def create_json_dataset(marketing_dialogues: List[str], 
                       real_dialogues: List[str],
                       output_path: str):
    """
    创建JSON格式数据集
    结构化存储，便于分析
    """
    dataset = {
        'metadata': {
            'total_count': len(marketing_dialogues) + len(real_dialogues),
            'marketing_count': len(marketing_dialogues),
            'real_count': len(real_dialogues),
            'format': 'dialogue',
            'placeholders': {
                'A': 'product',
                'B': 'brand',
                'C': 'activity',
                'D': 'platform'
            }
        },
        'data': []
    }
    
    # 添加营销对话
    for i, dialogue in enumerate(marketing_dialogues, 1):
        dataset['data'].append({
            'id': f'marketing_{i}',
            'dialogue': dialogue,
            'type': 'marketing',
            'source': 'synthetic',
            'label': 0
        })
    
    # 添加真实评论
    for i, dialogue in enumerate(real_dialogues, 1):
        dataset['data'].append({
            'id': f'real_{i}',
            'dialogue': dialogue,
            'type': 'user_comment',
            'source': 'real',
            'label': 1
        })
    
    with open(output_path, 'w', encoding='utf-8') as f:
        json.dump(dataset, f, ensure_ascii=False, indent=2)
    
    print(f"✅ JSON数据集已保存到: {output_path}")


def create_mixed_txt_dataset(marketing_dialogues: List[str], 
                             real_dialogues: List[str],
                             output_path: str):
    """
    创建纯文本格式数据集
    简单混合，用于直接阅读或简单训练
    """
    with open(output_path, 'w', encoding='utf-8') as f:
        f.write("营销对话 + 真实评论 混合数据集\n")
        f.write("=" * 70 + "\n\n")
        
        # 写入营销对话部分
        f.write("【第一部分：营销对话数据】\n")
        f.write(f"共 {len(marketing_dialogues)} 组\n")
        f.write("-" * 70 + "\n\n")
        
        for i, dialogue in enumerate(marketing_dialogues, 1):
            f.write(f"对话 {i} [营销]\n")
            f.write(dialogue + "\n\n")
        
        # 写入真实评论部分
        f.write("\n" + "=" * 70 + "\n")
        f.write("【第二部分：真实评论对话】\n")
        f.write(f"共 {len(real_dialogues)} 组\n")
        f.write("-" * 70 + "\n\n")
        
        for i, dialogue in enumerate(real_dialogues, 1):
            f.write(f"对话 {i} [真实]\n")
            f.write(dialogue + "\n\n")
    
    print(f"✅ 混合文本数据集已保存到: {output_path}")


def main():
    print("=" * 70)
    print("数据合并工具：营销对话 + 真实评论")
    print("=" * 70)
    
    # 加载数据
    print("\n📖 加载数据...")
    marketing_path = PROJECT_ROOT / "data" / "marketing" / "营销对话数据.docx"
    real_path = PROJECT_ROOT / "data" / "comment_to_dialogue" / "dialogues_enhanced.txt"
    
    marketing_dialogues = extract_marketing_dialogues(str(marketing_path))
    real_dialogues = load_real_dialogues(real_path)
    
    print(f"   营销对话: {len(marketing_dialogues)} 组")
    print(f"   真实评论: {len(real_dialogues)} 组")
    
    # 创建多种格式的数据集
    print("\n🔄 生成数据集...")
    output_dir = PROJECT_ROOT / "data"
    output_dir.mkdir(parents=True, exist_ok=True)
    
    # 1. JSONL格式（推荐用于大模型训练）
    create_jsonl_dataset(
        marketing_dialogues, 
        real_dialogues,
        output_dir / "merged_dataset.jsonl"
    )
    
    print()
    
    # 2. JSON格式（便于分析和查看）
    create_json_dataset(
        marketing_dialogues,
        real_dialogues,
        output_dir / "merged_dataset.json"
    )
    
    print()
    
    # 3. TXT格式（便于阅读）
    create_mixed_txt_dataset(
        marketing_dialogues,
        real_dialogues,
        output_dir / "merged_dataset.txt"
    )
    
    print("\n" + "=" * 70)
    print("✨ 数据合并完成！")
    print("=" * 70)
    print("\n生成的文件：")
    print("  1. merged_dataset.jsonl  - 用于模型训练（推荐）")
    print("  2. merged_dataset.json   - 用于数据分析")
    print("  3. merged_dataset.txt    - 用于直接阅读")
    print("\n使用建议：")
    print("  • 模型训练：使用 .jsonl 格式")
    print("  • 数据分析：使用 .json 格式")
    print("  • 人工审核：使用 .txt 格式")


if __name__ == '__main__':
    main()
